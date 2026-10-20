"""Profile read + check + one bulk fix on a 600-row file. Run in odoo-bin shell."""
import base64
import cProfile
import io
import pstats
import time

from openpyxl import load_workbook

N = 600
SRC = r"C:/Users/Lenovo/Documents/imsma-importer/samples/eore_sep26/DMA EORE Template English-Sep 2026.xlsm"
wb = load_workbook(SRC, keep_vba=False)
ws = wb["EORE"]
rows = [[c.value for c in ws[r]] for r in range(6, 10)]
for i in range(N):
    vals = list(rows[i % 4])
    vals[44] = f"{vals[44] or 'R'}-{i}"
    for j, v in enumerate(vals, start=1):
        ws.cell(6 + i, j).value = v
buf = io.BytesIO()
wb.save(buf)
data = buf.getvalue()
batch = env["dma.fr.batch"].create({
    "template_id": env.ref("dma_field_reports.template_eore").id,
    "office_id": env.ref("dma_field_reports.office_dma_hq").id,
    "source_file": base64.b64encode(data), "source_filename": f"big_{N}.xlsx"})

from odoo.addons.dma_field_reports.services import xlsx  # noqa: E402

t = time.time()
xlsx.read_sheet(data, "EORE", 4, 2, 6)
print(f"openpyxl read only: {time.time() - t:.1f}s")


def prof(label, fn):
    p = cProfile.Profile()
    t = time.time()
    p.enable()
    fn()
    p.disable()
    print(f"\n===== {label}: {time.time() - t:.1f}s")
    pstats.Stats(p).sort_stats("cumulative").print_stats(18)


prof("read", batch._read_file)
prof("check", lambda: (batch._run_check(), batch._update_state()))
cell = batch.cell_ids.filtered(lambda c: c.code == "required" and c.field_id.name == "MAC owner")[:1]
w = env["dma.fr.fix.wizard"].create({"cell_id": cell.id, "option_id": env["dma.fr.ref.option"].search(
    [("value", "=", "RMAC-N"), ("cdf_ref_id.cdf_id", "=", "c0a7-f58e-1773eacedb7-ef3a9ab4-10-dff")]).id})
prof("fix 300 rows", w.action_apply)
env.cr.rollback()
