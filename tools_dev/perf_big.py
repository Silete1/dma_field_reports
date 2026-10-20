"""Timing on a big file: the real sample's 4 rows repeated to N rows. Run in odoo-bin shell."""
import base64
import io
import time

from openpyxl import load_workbook

N = 600
SRC = r"C:/Users/Lenovo/Documents/imsma-importer/samples/eore_sep26/DMA EORE Template English-Sep 2026.xlsm"
wb = load_workbook(SRC, keep_vba=False)
ws = wb["EORE"]
rows = [[c.value for c in ws[r]] for r in range(6, 10)]
for i in range(N):
    vals = list(rows[i % 4])
    vals[44] = f"{vals[44] or 'R'}-{i}"  # AS: report id unique per row
    for j, v in enumerate(vals, start=1):
        ws.cell(6 + i, j).value = v
buf = io.BytesIO()
wb.save(buf)
batch = env["dma.fr.batch"].create({
    "template_id": env.ref("dma_field_reports.template_eore").id,
    "office_id": env.ref("dma_field_reports.office_dma_hq").id,
    "source_file": base64.b64encode(buf.getvalue()), "source_filename": f"big_{N}.xlsx"})
t = time.time()
batch._read_file()
t1 = time.time()
batch._run_check()
batch._update_state()
t2 = time.time()
print(f"{N} rows: read {t1 - t:.1f}s, check {t2 - t1:.1f}s; must-fix={batch.error_count} cells={len(batch.cell_ids)}")
from collections import Counter
print(Counter((c.code, c.raw) for c in batch.cell_ids if c.state == "error").most_common(6))
cell = batch.cell_ids.filtered(lambda c: c.code == "required" and c.field_id.name == "MAC owner")[:1]
w = env["dma.fr.fix.wizard"].create({"cell_id": cell.id, "option_id": env["dma.fr.ref.option"].search(
    [("value", "=", "RMAC-N"), ("cdf_ref_id.cdf_id", "=", "c0a7-f58e-1773eacedb7-ef3a9ab4-10-dff")]).id})
t3 = time.time()
w.action_apply()
print(f"one fix on {w.same_count} rows: {time.time() - t3:.1f}s; must-fix now {batch.error_count}")
env.cr.rollback()
