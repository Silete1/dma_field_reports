"""A batch = one uploaded operator file, its rows (one report each) and cells (one value each)."""
import base64
import io
import urllib.parse

from markupsafe import Markup
from psycopg2.extras import execute_values

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from ..services import text, xlsx

BATCH_STATES = [
    ("draft", "Uploaded"),
    ("checked", "Problems to fix"),
    ("locations", "New locations in IMSMA"),
    ("ready", "Ready to export"),
    ("exported", "Exported"),
    ("done", "In IMSMA"),
    ("cancel", "Cancelled"),
]
# cells kept even when empty: row-level checks flag them (place, coordinates, attachments)
ALWAYS_KEPT_ROLES = ("province", "district", "subdistrict", "city", "location_name", "latitude", "longitude", "mgrs",
                     "attachments", "org_report_id")
CELL_STATES = [("ok", "OK"), ("empty", "Empty"), ("info", "Note"), ("warning", "Check"), ("error", "Must fix")]


class DmaFrOffice(models.Model):
    _name = "dma.fr.office"
    _description = "Office that imports field data"
    _order = "sequence, name"

    name = fields.Char(required=True, translate=True)
    code = fields.Char(required=True, help="Goes into IMSMA ids: Education-<code>-123, LOC-<code>-45.")
    sequence = fields.Integer(default=10)
    active = fields.Boolean(default=True)

    _code_uniq = models.Constraint("unique(code)", "Office codes must be unique.")


class DmaFrCounter(models.Model):
    _name = "dma.fr.counter"
    _description = "IMSMA id counter used by exports"
    _rec_name = "id_type"

    id_type = fields.Char("IMSMA id type", required=True, help="IMSMA localidtype name, e.g. Education, LOC.")
    last_number = fields.Integer("Last number given by Odoo", readonly=True)
    note = fields.Char()

    _uniq = models.Constraint("unique(id_type)", "One counter per id type.")

    @api.model
    def allocate(self, id_type, count, prefixes):
        """Reserve `count` numbers for `id_type`. Never below the numbers IMSMA already uses.

        prefixes: local-id prefixes that share this IMSMA counter (e.g. 'Education-CI-', 'FR-MRE-SI-').
        Returns the first number. Locks the counter row, so two exports never get the same numbers.
        """
        counter = self.search([("id_type", "=", id_type)], limit=1) or self.create({"id_type": id_type})
        self.env.cr.execute("SELECT last_number FROM dma_fr_counter WHERE id = %s FOR UPDATE", [counter.id])
        last = self.env.cr.fetchone()[0] or 0
        Usage = self.env["dma.fr.ref.idusage"]
        floor = 0
        for u in Usage.search([]):
            if (u.source == "counter" and u.prefix == id_type) or (u.source == "usage" and u.prefix in prefixes):
                floor = max(floor, u.max_number)
        first = max(last, floor) + 1
        counter.last_number = first + count - 1
        return first


class DmaFrBatch(models.Model):
    _name = "dma.fr.batch"
    _description = "Field data batch"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "id desc"

    name = fields.Char(required=True, readonly=True, copy=False, default=lambda s: _("New"))
    template_id = fields.Many2one("dma.fr.template", "Form", required=True, tracking=True,
                                  default=lambda s: s.env["dma.fr.template"].search([("kind", "=", "eore")], limit=1))
    kind = fields.Selection(related="template_id.kind")
    template_file = fields.Binary(related="template_id.operator_file", string="Blank file for operators")
    template_filename = fields.Char(related="template_id.operator_filename")
    office_id = fields.Many2one("dma.fr.office", "Office", required=True, tracking=True,
                                default=lambda s: s.env["dma.fr.office"].search([], limit=1))
    source_file = fields.Binary("Operator file", attachment=True, copy=False)
    source_filename = fields.Char()
    read_filename = fields.Char(readonly=True, copy=False, help="File the rows were read from.")
    attachment_ids = fields.Many2many("ir.attachment", "dma_fr_batch_attachment_rel", "batch_id", "attachment_id",
                                      string="Attachments (reports, photos)", copy=False)
    state = fields.Selection(BATCH_STATES, default="draft", required=True, tracking=True, copy=False)
    user_id = fields.Many2one("res.users", "Responsible", default=lambda s: s.env.user, tracking=True)
    row_ids = fields.One2many("dma.fr.row", "batch_id", "Reports")
    cell_ids = fields.One2many("dma.fr.cell", "batch_id", "Values")
    location_request_ids = fields.One2many("dma.fr.location.request", "batch_id", "New locations")
    export_ids = fields.One2many("dma.fr.export", "batch_id", "Exports")
    source_id = fields.Many2one("dma.fr.source", "IMSMA data used", readonly=True, copy=False)
    source_stamp = fields.Char(readonly=True, copy=False)
    checked_on = fields.Datetime(readonly=True, copy=False)
    column_report = fields.Html(readonly=True, copy=False, sanitize=False)
    row_count = fields.Integer(compute="_compute_counts", store=True)
    error_count = fields.Integer("Must fix", compute="_compute_counts", store=True)
    warning_count = fields.Integer("To check", compute="_compute_counts", store=True)
    ok_count = fields.Integer("Ready", compute="_compute_counts", store=True)
    excluded_count = fields.Integer("Left out", compute="_compute_counts", store=True)
    location_count = fields.Integer(compute="_compute_location_count")
    locations_to_export = fields.Integer(compute="_compute_location_count")
    locations_waiting = fields.Integer(compute="_compute_location_count")
    next_step = fields.Html(compute="_compute_next_step", sanitize=False)
    note = fields.Text()

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get("name", _("New")) == _("New"):
                vals["name"] = self.env["ir.sequence"].next_by_code("dma.fr.batch") or _("New")
        return super().create(vals_list)

    @api.depends("row_ids.state", "row_ids.excluded")
    def _compute_counts(self):
        for rec in self:
            rows = rec.row_ids
            rec.row_count = len(rows)
            rec.excluded_count = len(rows.filtered("excluded"))
            live = rows.filtered(lambda r: not r.excluded)
            rec.error_count = len(live.filtered(lambda r: r.state == "error"))
            rec.warning_count = len(live.filtered(lambda r: r.state == "warning"))
            rec.ok_count = len(live.filtered(lambda r: r.state in ("ok", "info")))

    @api.depends("location_request_ids.state", "location_request_ids.row_ids.excluded")
    def _compute_location_count(self):
        for rec in self:
            needed = rec.location_request_ids.filtered(lambda l: l.row_ids.filtered(lambda r: not r.excluded))
            rec.location_count = len(rec.location_request_ids)
            rec.locations_to_export = len(needed.filtered(lambda l: l.state == "draft"))
            rec.locations_waiting = len(needed.filtered(lambda l: l.state == "exported"))

    @api.depends("state", "error_count", "warning_count", "row_count", "location_request_ids.state", "source_file")
    def _compute_next_step(self):
        for rec in self:
            rec.next_step = rec._next_step_html()

    def _next_step_html(self):
        self.ensure_one()

        def box(kind, title, body):
            return Markup('<div class="alert alert-%s mb-0" role="status"><strong>%s</strong><div>%s</div></div>') % (
                kind, title, body)

        if self.state == "cancel":
            return box("secondary", _("Cancelled"), _("Nothing from this batch will be exported."))
        if not self.row_count:
            if not self.source_file:
                return box("info", _("Step 1 of 4: upload"), _(
                    "Choose the form and office, upload the operator's Excel file and drop the attachments "
                    "(Word reports, photos) in Attachments. Then press Check."))
            return box("info", _("Step 1 of 4: check"), _("Press Check to read the file."))
        if self.state in ("draft", "checked") and self.error_count:
            return box("danger", _("Step 2 of 4: fix %s report(s)", self.error_count), _(
                "Open Fix problems. Problems are grouped by column and value: fixing one fixes every row with the "
                "same value. Use Leave out for a report that should not go to IMSMA."))
        pending = self.location_request_ids.filtered(lambda l: l.state != "confirmed")
        if pending and self.state in ("draft", "checked", "ready", "locations"):
            if any(l.state == "draft" for l in pending):
                return box("warning", _("Step 3 of 4: %s new location(s)", len(pending)), _(
                    "These places have no IMSMA location yet. Press Export new locations, import that zip in the "
                    "IMSMA Workbench, approve the locations, then press Locations approved."))
            return box("warning", _("Step 3 of 4: waiting for IMSMA"), _(
                "Import the locations zip in IMSMA and approve the location reports, then press Locations approved."))
        if self.state in ("draft", "checked", "ready"):
            extra = _(" %s report(s) have notes to check; they can still be exported.", self.warning_count) \
                if self.warning_count else ""
            return box("success", _("Step 4 of 4: export"), _("Everything required is in place.") + extra + " " + _(
                "Press Export to IMSMA and import the zip in the IMSMA Workbench."))
        if self.state == "exported":
            return box("success", _("Exported"), _(
                "Import the zip in the IMSMA Workbench and approve the reports. Odoo marks them as in IMSMA when the "
                "IMSMA data is next loaded."))
        return box("success", _("In IMSMA"), _("Every report of this batch was found in IMSMA."))

    # ---------------------------------------------------------------- reading
    def _read_file(self):
        """(Re)build rows and cells from the uploaded file."""
        self.ensure_one()
        if not self.source_file:
            raise UserError(_("Upload the operator's Excel file first."))
        t = self.template_id
        data = base64.b64decode(self.source_file)
        try:
            sheet = xlsx.read_sheet(data, t.sheet_name, t.header_row, t.header_rows, t.first_data_row)
        except KeyError as e:
            raise UserError(str(e.args[0])) from e
        except Exception as e:  # noqa: BLE001 - any unreadable file
            raise UserError(_("This file could not be read as Excel (%s).", e)) from e
        self.row_ids.unlink()
        self.location_request_ids.unlink()
        col_for, missing, used = {}, [], set()
        for f in t.field_ids:
            cols = [c for c in sheet.find(f.header_list()) if c.index not in used]
            if cols:
                col_for[f.id] = cols[0]
                used.add(cols[0].index)
            elif f.role != "ignore":
                missing.append(f)
        unknown = [c for c in sheet.columns if c.index not in used]
        rows = self.env["dma.fr.row"].create([{"batch_id": self.id, "excel_row": n} for n, _v in sheet.rows])
        # one value per row and column: tens of thousands for a big file, so one bulk insert
        # (the ORM would take minutes on the 1-vCPU server); stored related fields are filled here too
        uid, now = self.env.uid, fields.Datetime.now()
        data = []
        for row, (number, values) in zip(rows, sheet.rows):
            for f in t.field_ids:
                col = col_for.get(f.id)
                raw = values[col.index] if col and col.index < len(values) else None
                if raw in (None, "") and not f.required and f.role not in ALWAYS_KEPT_ROLES:
                    continue  # an empty optional value: nothing to check, show or send
                data.append((row.id, self.id, number, False, f.id, f.sequence, 0, col.letter if col else "",
                             _raw_text(raw), raw if hasattr(raw, "year") else None, "ok", uid, now, uid, now))
        if data:
            execute_values(self.env.cr._obj, """
                INSERT INTO dma_fr_cell (row_id, batch_id, excel_row, row_excluded, field_id, sequence, severity,
                                         "column", raw, raw_date, state, create_uid, create_date, write_uid, write_date)
                VALUES %s""", data, page_size=2000)
        self.env.invalidate_all()
        rows.modified(["cell_ids"])
        self.env.flush_all()
        self.read_filename = self.source_filename
        report = Markup("<p>%s</p>") % _("%(rows)s report(s) read from sheet '%(sheet)s'.",
                                          rows=len(sheet.rows), sheet=sheet.name)
        if missing:
            report += Markup("<p class='text-danger'>%s %s</p>") % (
                _("Columns not found in the file:"), ", ".join(f.name for f in missing))
        if unknown:
            report += Markup("<p class='text-muted'>%s %s</p>") % (
                _("Columns in the file that are not used:"), ", ".join(f"{c.letter} {c.path}" for c in unknown))
        self.column_report = report

    # ---------------------------------------------------------------- actions
    def action_check(self):
        for rec in self:
            if rec.state in ("exported", "done", "cancel"):
                raise UserError(_("This batch is already exported. Use Re-open to change it."))
            source = self.env["dma.fr.source"].current()
            if not source:
                raise UserError(_("No IMSMA data is loaded. A manager must load it first "
                                  "(Field Reports > Configuration > IMSMA data source)."))
            if not rec.row_ids or rec.read_filename != rec.source_filename:
                rec._read_file()
            rec._run_check()
            rec.write({"source_id": source.id, "source_stamp": source.stamp, "checked_on": fields.Datetime.now()})
            rec._update_state()
        return True

    def _update_state(self):
        for rec in self:
            if rec.state in ("exported", "done", "cancel"):
                continue
            if rec.error_count:
                rec.state = "checked"
            elif rec.locations_to_export or rec.locations_waiting:
                rec.state = "locations"  # new IMSMA locations come first (step 3)
            else:
                rec.state = "ready" if rec.row_count else "draft"

    def action_open_problems(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "name": _("Fix problems: %s", self.name),
            "res_model": "dma.fr.cell",
            "view_mode": "list,form",
            "views": [(self.env.ref("dma_field_reports.dma_fr_cell_view_problem_list").id, "list"), (False, "form")],
            "domain": [("batch_id", "=", self.id), ("state", "in", ("error", "warning")), ("row_excluded", "=", False)],
            "context": {"group_by": ["field_id"]},
            "target": "current",
        }

    def action_open_rows(self):
        self.ensure_one()
        return {"type": "ir.actions.act_window", "name": _("Reports: %s", self.name), "res_model": "dma.fr.row",
                "view_mode": "list,form", "domain": [("batch_id", "=", self.id)],
                "context": {"default_batch_id": self.id}}

    def action_open_locations(self):
        self.ensure_one()
        return {"type": "ir.actions.act_window", "name": _("New locations: %s", self.name),
                "res_model": "dma.fr.location.request", "view_mode": "list,form",
                "domain": [("batch_id", "=", self.id)]}

    def action_problem_report(self):
        """Excel list of the problems, to send back to the operator (in the user's language)."""
        self.ensure_one()
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        cells = self.cell_ids.filtered(lambda c: c.state in ("error", "warning") and not c.row_excluded).sorted(
            lambda c: (c.excel_row, c.sequence))
        wb = Workbook()
        ws = wb.active
        ws.title = _("Problems")[:31]
        lang = self.env.lang or "en_US"
        ws.sheet_view.rightToLeft = lang.startswith("ar")
        ws.append([_("Batch"), self.name, _("File"), self.source_filename or "", _("Checked on"),
                   fields.Datetime.to_string(self.checked_on) or ""])
        ws.append([])
        head = [_("Row"), _("Column"), _("Field"), _("In the file"), _("Problem"), _("Closest IMSMA values"),
                _("Severity")]
        ws.append(head)
        bold = Font(bold=True)
        for c in ws[3]:
            c.font = bold
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        sev = dict(self.env["dma.fr.cell"]._fields["state"]._description_selection(self.env))
        red, amber = PatternFill("solid", fgColor="F8D7DA"), PatternFill("solid", fgColor="FFF3CD")
        for cell in cells:
            ws.append([cell.excel_row, cell.column or "", cell.field_id.name, cell.raw or "", cell.message or "",
                       (cell.suggestion or "").replace(" | ", ", "), sev.get(cell.state, cell.state)])
            ws.cell(ws.max_row, 7).fill = red if cell.state == "error" else amber
        for col, width in zip("ABCDEFG", (7, 8, 26, 28, 70, 34, 12)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=4):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A4"
        buf = io.BytesIO()
        wb.save(buf)
        name = f"{self.name.replace('/', '-')}_{_('problems')}.xlsx"
        att = self.env["ir.attachment"].create({
            "name": name, "datas": base64.b64encode(buf.getvalue()), "res_model": self._name, "res_id": self.id,
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
        return {"type": "ir.actions.act_url", "target": "self",
                "url": f"/web/content/{att.id}?download=true"}

    def action_reopen(self):
        for rec in self:
            rec.state = "checked"
            rec.message_post(body=_("Re-opened. A new export keeps the IMSMA ids already given, so IMSMA updates "
                                    "the same reports."))
            rec.action_check()

    def action_cancel(self):
        self.write({"state": "cancel"})

    def action_reset(self):
        self.write({"state": "draft"})


def _raw_text(value):
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class DmaFrRow(models.Model):
    _name = "dma.fr.row"
    _description = "Field report (one row of the file)"
    _order = "batch_id, excel_row"
    _rec_name = "summary"

    batch_id = fields.Many2one("dma.fr.batch", required=True, index=True, ondelete="cascade")
    excel_row = fields.Integer("Row", readonly=True)
    cell_ids = fields.One2many("dma.fr.cell", "row_id", "Values")
    state = fields.Selection(CELL_STATES, compute="_compute_state", store=True)
    problem_count = fields.Integer(compute="_compute_state", store=True)
    excluded = fields.Boolean("Left out", help="Not exported to IMSMA.")
    summary = fields.Char(compute="_compute_summary", store=True)
    # resolved values
    org_id = fields.Many2one("dma.fr.ref.org", "Organisation", readonly=True)
    report_date = fields.Date(readonly=True)
    place_id = fields.Many2one("dma.fr.ref.place", "Place (IMSMA gazetteer)")
    place_manual = fields.Boolean("Place chosen by hand")
    area_id = fields.Many2one("dma.fr.ref.place", "Area found", readonly=True,
                              help="Deepest level of the address that matched the gazetteer (e.g. the district).")
    location_manual_id = fields.Many2one("dma.fr.ref.location", "IMSMA location chosen by hand")
    location_id = fields.Many2one("dma.fr.ref.location", "IMSMA location", readonly=True)
    location_request_id = fields.Many2one("dma.fr.location.request", "New location", readonly=True)
    location_label = fields.Char("Location", compute="_compute_location_label")
    latitude = fields.Float(digits=(10, 7), readonly=True)
    longitude = fields.Float(digits=(10, 7), readonly=True)
    coord_note = fields.Char(readonly=True)
    map_url = fields.Char(compute="_compute_map_url")
    attachment_ids = fields.Many2many("ir.attachment", "dma_fr_row_attachment_rel", "row_id", "attachment_id",
                                      string="Attachments")
    # given at export, then frozen
    object_localid = fields.Char("IMSMA id", readonly=True, copy=False)
    report_localid = fields.Char("Field report id", readonly=True, copy=False)
    object_guid = fields.Char(readonly=True, copy=False)
    fieldreport_guid = fields.Char(readonly=True, copy=False)
    imsma_state = fields.Selection([("not_sent", "Not exported"), ("sent", "Exported"), ("found", "In IMSMA")], string="IMSMA",
                                   default="not_sent", readonly=True, copy=False)
    imsma_status = fields.Char("IMSMA status", readonly=True, copy=False)

    @api.depends("cell_ids.state", "excluded")
    def _compute_state(self):
        rank = {"ok": 0, "empty": 0, "info": 1, "warning": 2, "error": 3}
        for rec in self:
            states = rec.cell_ids.mapped("state")
            worst = max(states, key=lambda s: rank.get(s, 0)) if states else "ok"
            rec.state = "ok" if worst == "empty" else worst
            rec.problem_count = len([s for s in states if s in ("warning", "error")])

    @api.depends("cell_ids.value", "cell_ids.raw", "excel_row")
    def _compute_summary(self):
        for rec in self:
            def val(role):
                c = rec.cell_ids.filtered(lambda c: c.field_id.role == role)[:1]
                return c.value or c.raw if c else ""
            parts = [val("location_name") or val("city"), val("report_date")]
            rec.summary = _("Row %s", rec.excel_row) + ": " + " · ".join(p for p in parts if p)

    @api.depends("location_id", "location_request_id")
    def _compute_location_label(self):
        for rec in self:
            if rec.location_id:
                rec.location_label = rec.location_id.display_name
            elif rec.location_request_id:
                rec.location_label = _("New: %s", rec.location_request_id.name)
            else:
                rec.location_label = False

    @api.depends("latitude", "longitude")
    def _compute_map_url(self):
        for rec in self:
            if rec.latitude or rec.longitude:
                rec.map_url = "https://www.openstreetmap.org/?" + urllib.parse.urlencode(
                    {"mlat": f"{rec.latitude:.6f}", "mlon": f"{rec.longitude:.6f}"}) + \
                    f"#map=14/{rec.latitude:.5f}/{rec.longitude:.5f}"
            else:
                rec.map_url = False

    def action_exclude(self):
        self.write({"excluded": True})
        self.batch_id._update_state()

    def action_include(self):
        self.write({"excluded": False})
        self.batch_id._run_check(self)
        self.batch_id._update_state()

    @api.model
    def _reconcile_with_imsma(self, report_rows):
        """Mark exported rows whose field report is now in IMSMA. Returns stats."""
        sent = self.search([("imsma_state", "in", ("sent", "found")), ("fieldreport_guid", "!=", False)])
        if not sent:
            return {}
        by_guid = {r["guid"]: r for r in report_rows}
        by_localid = {r["localid"]: r for r in report_rows}
        found = 0
        for row in sent:
            hit = by_guid.get(row.fieldreport_guid) or by_localid.get(row.report_localid)
            if hit:
                row.write({"imsma_state": "found", "imsma_status": hit.get("status") or ""})
                found += 1
        for batch in sent.batch_id:
            rows = batch.row_ids.filtered(lambda r: not r.excluded)
            if rows and all(r.imsma_state == "found" for r in rows) and batch.state == "exported":
                batch.state = "done"
                batch.message_post(body=_("All reports of this batch are in IMSMA."))
        for req in self.env["dma.fr.location.request"].search([("state", "=", "exported")]):
            if req.fieldreport_guid in by_guid or req.report_localid in by_localid:
                req.imsma_seen = True
        return {"exported reports found in IMSMA": found}


class DmaFrCell(models.Model):
    _name = "dma.fr.cell"
    _description = "Value of a field report"
    _order = "row_id, sequence, id"

    row_id = fields.Many2one("dma.fr.row", required=True, index=True, ondelete="cascade")
    batch_id = fields.Many2one(related="row_id.batch_id", store=True, index=True)
    row_excluded = fields.Boolean(related="row_id.excluded", store=True)
    excel_row = fields.Integer(related="row_id.excel_row", string="Excel row", store=True, aggregator=False)
    field_id = fields.Many2one("dma.fr.template.field", "Column", required=True, ondelete="cascade")
    sequence = fields.Integer(related="field_id.sequence", store=True, aggregator=False)
    severity = fields.Integer(compute="_compute_severity", store=True, aggregator=False)
    column = fields.Char("Col", readonly=True)
    raw = fields.Char("In the file", readonly=True)
    raw_date = fields.Date(readonly=True)
    override = fields.Char("Corrected value",
                           help="Type the right value here; the row is checked again when you save.")
    value = fields.Char("Checked value", readonly=True,
                        help="Stored only when it differs from the file (a translated list value, a corrected date...).")
    sent_value = fields.Char("Sent to IMSMA", compute="_compute_sent_value")
    state = fields.Selection(CELL_STATES, default="ok", readonly=True, index=True)
    code = fields.Char(readonly=True)
    message = fields.Char(readonly=True)
    suggestion = fields.Char(readonly=True, help="Best guesses, best first, separated by ' | '.")
    same_count = fields.Integer("Rows with this problem", compute="_compute_same_count")

    @api.depends("value", "raw", "state")
    def _compute_sent_value(self):
        for rec in self:
            if rec.value:
                rec.sent_value = rec.value
            else:
                rec.sent_value = rec.raw if rec.state in ("ok", "info", "warning") else False

    @api.depends("state")
    def _compute_severity(self):
        rank = {"error": 3, "warning": 2, "info": 1}
        for rec in self:
            rec.severity = rank.get(rec.state, 0)

    def _compute_same_count(self):
        for rec in self:
            rec.same_count = self.search_count([("batch_id", "=", rec.batch_id.id), ("field_id", "=", rec.field_id.id),
                                                ("raw", "=", rec.raw), ("code", "=", rec.code)])

    def write(self, vals):
        res = super().write(vals)
        if "override" in vals and not self.env.context.get("dma_fr_no_recheck"):
            for batch in self.batch_id:
                batch._run_check(self.row_id.filtered(lambda r: r.batch_id == batch))
                batch._update_state()
        return res

    def action_fix(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "name": _("Fix: %s", self.field_id.name),
            "res_model": "dma.fr.fix.wizard",
            "view_mode": "form",
            "target": "new",
            "context": {"default_cell_id": self.id},
        }

    def action_accept(self):
        """Accept a warning as it is (e.g. a long comment, a missing attachment)."""
        self.with_context(dma_fr_no_recheck=True).write({"state": "ok", "message": _("Accepted by %s", self.env.user.name)})


class DmaFrLocationRequest(models.Model):
    _name = "dma.fr.location.request"
    _description = "Location to create in IMSMA"
    _order = "batch_id, name"

    batch_id = fields.Many2one("dma.fr.batch", required=True, index=True, ondelete="cascade")
    name = fields.Char("Location name", required=True)
    place_id = fields.Many2one("dma.fr.ref.place", "Gazetteer place", required=True)
    row_ids = fields.One2many("dma.fr.row", "location_request_id", "Reports")
    row_count = fields.Integer(compute="_compute_row_count")
    state = fields.Selection([("draft", "To export"), ("exported", "Exported, waiting for approval in IMSMA"),
                              ("confirmed", "Approved in IMSMA")], default="draft", required=True)
    localid = fields.Char("IMSMA id", readonly=True)
    guid = fields.Char(readonly=True)
    report_localid = fields.Char(readonly=True)
    fieldreport_guid = fields.Char(readonly=True)
    imsma_seen = fields.Boolean("Seen in IMSMA data", readonly=True)
    ref_location_id = fields.Many2one("dma.fr.ref.location", readonly=True)

    def _compute_row_count(self):
        for rec in self:
            rec.row_count = len(rec.row_ids)


class DmaFrExport(models.Model):
    _name = "dma.fr.export"
    _description = "IMSMA zip produced from a batch"
    _order = "id desc"

    batch_id = fields.Many2one("dma.fr.batch", required=True, index=True, ondelete="cascade")
    kind = fields.Selection([("locations", "New locations"), ("reports", "Field reports")], required=True)
    file = fields.Binary(attachment=True, readonly=True)
    filename = fields.Char(readonly=True)
    report_count = fields.Integer(readonly=True)
    counter_note = fields.Text("Set these IMSMA counters", readonly=True)
    create_uid = fields.Many2one("res.users", "Exported by", readonly=True)
    create_date = fields.Datetime("Exported on", readonly=True)
