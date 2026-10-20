"""Arabic files, repeated reports across batches, the operator's problem list."""
import base64
import datetime
import io

from openpyxl import Workbook, load_workbook

from odoo.tests import TransactionCase, tagged

from .test_flow import ROWS, _package, _workbook

D = datetime.datetime


def _arabic_workbook():
    """The same EORE layout with Arabic headers (as in the DMA Arabic form) and Arabic answers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "EORE"
    top = ["الجهة الممولة", "تاريخ التقرير", "أعد التقرير من قبل", "الموقع الوظيفي", "الجهة المنفذة", "رقم أمر العمل",
           "الموقع", None, None, None, None, None, "تاريخ البدء الفعلي", "تاريخ الانتهاء الفعلي",
           "توعية المجتمعات المتأثرة من مخاطر", "نوع التوعية بمخاطر الذخائر والمواد المتفجرة",
           "عائدية مركز شؤون الألغام"]
    sub = [None] * 6 + ["المحافظة", "القضاء", "القرية / المدينة", "MGRS", "خطوط العرض", "خطوط الطول"] + [None] * 5
    for i, (t, s) in enumerate(zip(top, sub), start=1):
        ws.cell(4, i).value = t
        ws.cell(5, i).value = s
    row = ["PMWRA", D(2026, 9, 27), "عامر زياد", "قائد الفريق", "NPA", "CI-NPA-EORE-1-26", "Anbar", "Heet", "Heet",
           None, 33.6401, 42.8301, D(2026, 9, 6), D(2026, 9, 17), "مخلفات حربية، عبوات ناسفة", "عرض توضيحي",
           "المركز الاقليمي الشمالي لشؤون الالغام"]
    for c, v in enumerate(row, start=1):
        ws.cell(6, c).value = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@tagged("post_install", "-at_install", "dma_fr")
class TestMore(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env["dma.fr.source"].search([]).write({"is_current": False})
        # independent of what this database has learned from earlier fixes
        cls.env["dma.fr.value.map"].search([("origin", "=", "learned")]).unlink()
        cls.source = cls.env["dma.fr.source"].create({
            "name": "test package", "mode": "package", "package_file": base64.b64encode(_package())})
        cls.source.action_refresh()
        # options the Arabic answers translate to (the invented package only has a few)
        Cdf = cls.env["dma.fr.ref.cdf"].with_context(active_test=False)
        guid = "7f00-0001-13f7a2fa1c2-3963ba9f-39-b37"
        cdf = Cdf.search([("cdf_id", "=", guid)]) or Cdf.create(
            {"cdf_id": guid, "name": "MRE Type", "datatype": "SINGLE_SELECT", "entity": "MRE"})
        cdf.active = True
        Opt = cls.env["dma.fr.ref.option"].with_context(active_test=False)
        opt = Opt.search([("cdf_ref_id", "=", cdf.id), ("value", "=", "Presentation")]) or Opt.create(
            {"cdf_ref_id": cdf.id, "value": "Presentation", "value_norm": "presentation", "usage": 5})
        opt.active = True

    def _batch(self, data, name="eore.xlsx"):
        return self.env["dma.fr.batch"].create({
            "template_id": self.env.ref("dma_field_reports.template_eore").id,
            "office_id": self.env.ref("dma_field_reports.office_dma_hq").id,
            "source_file": base64.b64encode(data), "source_filename": name})

    def test_arabic_headers_and_answers(self):
        b = self._batch(_arabic_workbook(), "eore_ar.xlsx")
        b.action_check()
        self.assertEqual(b.row_count, 1)
        cells = {c.field_id.name: c for c in b.cell_ids}
        # every Arabic header found its column
        for name in ("Funded by", "Date of report", "Reported by", "Organisation", "Province", "City / town / village",
                     "Latitude", "Awareness type", "EORE type", "MAC owner"):
            self.assertIn(name, cells, name)
        self.assertEqual(cells["Awareness type"].sent_value, "ERW, IED")  # Arabic comma split, both kept
        self.assertEqual(cells["EORE type"].sent_value, "Presentation")
        self.assertEqual(cells["MAC owner"].sent_value, "RMAC-N")
        self.assertEqual(b.error_count, 0, b.cell_ids.filtered(lambda c: c.state == "error").mapped("message"))
        self.assertEqual(b.row_ids.location_id.localid, "LOC-CI-100")  # the 'Heet' location under Heet city

    def test_report_already_exported_in_earlier_batch(self):
        first = self._batch(_workbook(ROWS[:1]))
        first.action_check()
        first.row_ids.write({"imsma_state": "sent", "object_localid": "Education-CI-9"})
        second = self._batch(_workbook(ROWS[:1]))
        second.action_check()
        cell = second.cell_ids.filtered(lambda c: c.field_id.role == "org_report_id")
        self.assertEqual(cell.code, "already_exported")
        self.assertIn(first.name, cell.message)

    def test_problem_list_for_the_operator(self):
        b = self._batch(_workbook(ROWS))
        b.action_check()
        action = b.action_problem_report()
        att = self.env["ir.attachment"].browse(int(action["url"].split("/web/content/")[1].split("?")[0]))
        ws = load_workbook(io.BytesIO(base64.b64decode(att.datas))).active
        problems = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
        self.assertEqual(len(problems), len(b.cell_ids.filtered(lambda c: c.state in ("error", "warning"))))
        self.assertIn("WRA", [p[3] for p in problems])
