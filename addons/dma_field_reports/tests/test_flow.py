"""Upload -> check -> fix -> new locations -> export, on invented IMSMA data (no real IMSMA data in git)."""
import base64
import datetime
import io
import zipfile

from lxml import etree
from openpyxl import Workbook

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged

from ..services import refdata

FUNDED = "c0a8-cc01-160e6ac65d2-2fbf5fc1-1d-ab3"
MAC = "c0a7-f58e-1773eacedb7-ef3a9ab4-10-dff"
AWARE = "c0a8-429c-1524042ff34-acafed6d-e89-45"


def _package():
    out = {name: [] for name in refdata.FILES}
    out["forms.csv"] = [
        {"guid": "form-eore", "name": "DMA Explosive Ordnance Risk Education Form", "version": "23",
         "published_on": "2023-02-03 10:00:00", "active": "true"},
        {"guid": "form-loc", "name": "DMA Location Form", "version": "20", "published_on": "2023-02-03 10:00:00",
         "active": "true"},
    ]
    out["cdfs.csv"] = [
        {"cdf_id": FUNDED, "name": "cdf_project_funded_by", "datatype": "SINGLE_SELECT", "entity": "MRE", "label": "", "active": "true"},
        {"cdf_id": MAC, "name": "mac_owner", "datatype": "SINGLE_SELECT", "entity": "MRE", "label": "", "active": "true"},
        {"cdf_id": AWARE, "name": "aw", "datatype": "MULTI_SELECT", "entity": "MRE", "label": "", "active": "true"},
    ]
    out["cdf_options.csv"] = [
        {"cdf_id": FUNDED, "value": "PMWRA", "usage": "100"}, {"cdf_id": FUNDED, "value": "Norway", "usage": "50"},
        {"cdf_id": MAC, "value": "RMAC-N", "usage": "10"}, {"cdf_id": MAC, "value": "RMAC-S", "usage": "5"},
        {"cdf_id": AWARE, "value": "IED", "usage": "9"}, {"cdf_id": AWARE, "value": "ERW", "usage": "9"},
    ]
    out["places.csv"] = [
        {"guid": "{IRQ}", "name": "Iraq", "level": "country", "parent_guid": "", "lat": "", "lon": ""},
        {"guid": "{P1}", "name": "Anbar", "level": "province", "parent_guid": "{IRQ}", "lat": "", "lon": ""},
        {"guid": "{D1}", "name": "Heet", "level": "district", "parent_guid": "{P1}", "lat": "", "lon": ""},
        {"guid": "{S1}", "name": "Heet", "level": "subdistrict", "parent_guid": "{D1}", "lat": "", "lon": ""},
        {"guid": "{C1}", "name": "Yardah", "level": "city", "parent_guid": "{S1}", "lat": "", "lon": ""},
        {"guid": "{C2}", "name": "Heet", "level": "city", "parent_guid": "{S1}", "lat": "", "lon": ""},
    ]
    out["locations.csv"] = [
        {"guid": "loc-1", "localid": "LOC-CI-100", "name": "Heet", "place_guid": "{C2}", "active": "true",
         "lat": "33.64", "lon": "42.83"},
        {"guid": "loc-2", "localid": "LOC-CI-101", "name": "قرية الكسارة", "place_guid": "{C2}", "active": "true",
         "lat": "33.7115", "lon": "42.7110"},
    ]
    out["organisations.csv"] = [{"guid": "{ORG-NPA}", "localid": "NPA", "name": "Norwegian People Aid", "active": "true"}]
    out["tasks.csv"] = [{"guid": "t1", "localid": "CI-NPA-EORE-1-26"}]
    out["id_usage.csv"] = [{"prefix": "Education-CI-", "max_number": "500"}, {"prefix": "LOC-CI-", "max_number": "100"}]
    buf = io.BytesIO()
    refdata.write_package(out, {"source": "backup", "dump_date": "2025-03-02"}, buf)
    return buf.getvalue()


def _workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "EORE"
    top = ["Funded by*", "Date of report*", "Reported by*", "Position*", "Organization*", "Task dossier ID*",
           "Location *", None, None, None, None, None, None, None, "Actual start date*", "Actual end date*",
           "Awareness type for impacted communities", "Org Internal Report ID", "MAC Owner*"]
    sub = [None] * 6 + ["Province", "District", "Subdistrict", "City/Town/Village", "Location name", "MGRS", "Lat ",
                        "Long"] + [None] * 5
    for i, (t, s) in enumerate(zip(top, sub), start=1):
        ws.cell(4, i).value = t
        ws.cell(5, i).value = s
    for r, values in enumerate(rows, start=6):
        for c, v in enumerate(values, start=1):
            ws.cell(r, c).value = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


D = datetime.datetime
ROWS = [
    # funded, date, by, position, org, dossier, prov, dist, sub, city, locname, mgrs, lat, lon, start, end, aware, reportid, mac
    ["WRA", D(2026, 9, 27), "Amer", "Team Leader", "NPA", "CI-NPA-EORE-1-26", "Anbar", "Heet", None, "Alkasarah", None,
     "38SKC8778732517", None, 42.70998, D(2026, 9, 6), D(2026, 9, 17), "IED,ERW", "ANB-1", "RMAC-N"],
    ["WRA", D(2026, 9, 27), "Amer", "Team Leader", "Norwegian People Aid", "CI-NPA-EORE-9-26", "Anbar", "Heet", None,
     "Yardah", None, None, 42.605468, 33.879166, D(2026, 9, 20), D(2026, 9, 20), "IED", "ANB-2", None],
]


@tagged("post_install", "-at_install", "dma_fr")
class TestFlow(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env["dma.fr.source"].search([]).write({"is_current": False})
        # independent of what this database has learned from earlier fixes
        cls.env["dma.fr.value.map"].search([("origin", "=", "learned")]).unlink()
        cls.source = cls.env["dma.fr.source"].create({
            "name": "test package", "mode": "package", "package_file": base64.b64encode(_package())})
        cls.source.action_refresh()
        cls.template = cls.env.ref("dma_field_reports.template_eore")
        cls.office = cls.env.ref("dma_field_reports.office_dma_hq")
        att = cls.env["ir.attachment"].create({"name": "ANB-1.docx", "datas": base64.b64encode(b"report one")})
        cls.batch = cls.env["dma.fr.batch"].create({
            "template_id": cls.template.id, "office_id": cls.office.id,
            "source_file": base64.b64encode(_workbook(ROWS)), "source_filename": "eore.xlsx",
            "attachment_ids": [(6, 0, att.ids)]})

    def cell(self, row, name):
        return self.batch.cell_ids.filtered(lambda c: c.excel_row == row and c.field_id.name == name)

    def fix(self, cell, **vals):
        w = self.env["dma.fr.fix.wizard"].create(dict({"cell_id": cell.id}, **vals))
        w.action_apply()

    def test_full_flow(self):
        b = self.batch
        b.action_check()
        self.assertEqual(b.row_count, 2)
        self.assertEqual(b.state, "checked")
        # every problem is listed at once, nothing aborted
        self.assertEqual(self.cell(6, "Funded by").code, "unknown_value")
        self.assertIn("PMWRA", self.cell(6, "Funded by").suggestion)
        self.assertEqual(self.cell(6, "City / town / village").code, "unknown_place")
        self.assertIn("قرية الكسارة", self.cell(6, "City / town / village").message)  # nearest IMSMA location
        self.assertEqual(self.cell(6, "Latitude").code, "from_mgrs")
        self.assertEqual(self.cell(7, "Latitude").code, "latlon_swapped")
        self.assertEqual(self.cell(7, "MAC owner").code, "required")
        self.assertEqual(self.cell(7, "Task dossier id").code, "unknown_task")  # warning only
        row6 = b.row_ids.filtered(lambda r: r.excel_row == 6)
        self.assertEqual(row6.attachment_ids.mapped("name"), ["ANB-1.docx"])
        with self.assertRaises(UserError):
            b.action_export()

        # one fix for both rows, remembered for next time
        option = self.env["dma.fr.ref.option"].search([("value", "=", "PMWRA"), ("cdf_ref_id.cdf_id", "=", FUNDED)])
        self.fix(self.cell(6, "Funded by"), option_id=option.id)
        self.assertFalse(self.cell(7, "Funded by").code)
        self.assertTrue(self.env["dma.fr.value.map"].search([("scope", "=", FUNDED), ("source_value", "=", "WRA")]))
        # village = existing IMSMA location nearby
        self.fix(self.cell(6, "City / town / village"), place_choice="existing",
                 location_pick_id=self.env["dma.fr.ref.location"].search([("localid", "=", "LOC-CI-101")]).id)
        self.fix(self.cell(7, "Latitude"))  # swap
        mac = self.env["dma.fr.ref.option"].search([("value", "=", "RMAC-N"), ("cdf_ref_id.cdf_id", "=", MAC)])
        self.fix(self.cell(7, "MAC owner"), option_id=mac.id)
        self.assertEqual(b.error_count, 0, b.cell_ids.filtered(lambda c: c.state == "error").mapped("message"))

        # Yardah has no IMSMA location: created first
        self.assertEqual(b.location_request_ids.mapped("name"), ["Yardah"])
        with self.assertRaises(UserError):
            b.action_export()
        b.action_export_locations()
        req = b.location_request_ids
        # numbers continue after IMSMA's highest (100 in the package) and never go back
        self.assertTrue(req.localid.startswith("LOC-CI-") and int(req.localid.rsplit("-", 1)[1]) > 100)
        self.assertEqual(b.state, "locations")
        b.action_confirm_locations()
        self.assertEqual(b.state, "ready")

        b.action_export()
        self.assertEqual(b.state, "exported")
        rows = b.row_ids.sorted("excel_row")
        numbers = [int(x.rsplit("-", 1)[1]) for x in rows.mapped("object_localid")]
        self.assertTrue(all(x.startswith("Education-CI-") for x in rows.mapped("object_localid")))
        self.assertGreater(numbers[0], 500)
        self.assertEqual(numbers[1], numbers[0] + 1)
        first = rows[0].report_localid
        export = b.export_ids.filtered(lambda e: e.kind == "reports")
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(export.file))) as z:
            names = z.namelist()
            self.assertIn(f"xml/{first}.xml", names)
            self.assertTrue(any(n.startswith("attachments/1.") and n.endswith(".docx") for n in names))
            root = etree.fromstring(z.read(f"xml/{first}.xml"))
        self.assertEqual(root.find(".//FieldReportInfo").get("reportGUID"), "form-eore")
        self.assertEqual(root.findtext(".//MREOrganisation/OrganisationID/OriginalID"), "NPA")
        self.assertEqual(root.findtext(".//Association/AssociationIdentification/OriginalID"), "LOC-CI-101")
        values = {c.findtext("CDF_ID"): [v.text for v in c.findall("Value")] for c in root.iter("CDF")}
        self.assertEqual(values["MRE.SINGLE_SELECT.cdf_project_funded_by_international"], ["PMWRA"])
        self.assertEqual(values["MRE.MULTI_SELECT.CDF-Awareness type for impacted communities from"], ["IED", "ERW"])
        self.assertAlmostEqual(float(root.findtext(".//XY/Y")), 33.7114, places=3)

        # the next IMSMA data shows the reports: batch is done
        rep = [{"guid": r.fieldreport_guid, "localid": r.report_localid, "status": "Approved"} for r in rows]
        self.env["dma.fr.row"]._reconcile_with_imsma(rep)
        self.assertEqual(b.state, "done")

    def test_reexport_keeps_ids(self):
        b = self.batch
        b.action_check()
        b.row_ids.filtered(lambda r: r.excel_row == 7).action_exclude()
        option = self.env["dma.fr.ref.option"].search([("value", "=", "PMWRA"), ("cdf_ref_id.cdf_id", "=", FUNDED)])
        self.fix(self.cell(6, "Funded by"), option_id=option.id)
        self.fix(self.cell(6, "City / town / village"), place_choice="existing",
                 location_pick_id=self.env["dma.fr.ref.location"].search([("localid", "=", "LOC-CI-101")]).id)
        b.action_export()
        first = b.row_ids.filtered(lambda r: r.excel_row == 6).object_localid
        b.action_reopen()
        b.action_export()
        self.assertEqual(b.row_ids.filtered(lambda r: r.excel_row == 6).object_localid, first)
        self.assertEqual(len(b.export_ids), 2)
