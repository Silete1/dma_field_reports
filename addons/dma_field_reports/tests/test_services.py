import datetime
import io

from lxml import etree
from openpyxl import Workbook

from odoo.tests import BaseCase, tagged

from ..services import geo, maxml, refdata, text, xlsx


@tagged("post_install", "-at_install", "dma_fr")
class TestText(BaseCase):
    def test_norm_latin_and_arabic(self):
        self.assertEqual(text.norm("Al-Kasarah "), text.norm("alkasarah"))
        self.assertEqual(text.norm("MAC Owner*"), "macowner")
        # alef with hamza / plain alef, ta marbuta / ha, diacritics
        self.assertEqual(text.norm("الإنبار"), text.norm("الانبار"))
        self.assertEqual(text.norm("مدرسة"), text.norm("مدرسه"))
        self.assertEqual(text.norm("عَرْض"), text.norm("عرض"))

    def test_split_multi_keeps_every_value(self):
        self.assertEqual(text.split_multi("IED,ERW"), ["IED", "ERW"])
        self.assertEqual(text.split_multi(" IED ، ERW ; Landmine"), ["IED", "ERW", "Landmine"])

    def test_best_matches(self):
        hits = text.best_matches("WRA", [(v, v) for v in ("PMWRA", "UN", "Norway")])
        self.assertEqual(hits[0][2], "PMWRA")


@tagged("post_install", "-at_install", "dma_fr")
class TestGeo(BaseCase):
    def test_mgrs_round_trip_matches_file(self):
        # real pair from an operator file: MGRS and lat/long agree to about a metre
        lat, lon = geo.parse_mgrs("38SKC7853351339")
        self.assertLess(geo.distance_km(lat, lon, 33.879166, 42.605468), 0.01)
        self.assertEqual(geo.to_mgrs(lat, lon), "38SKC7853351339")
        self.assertIsNotNone(geo.parse_mgrs("38S KC 78533 51339"))
        self.assertIsNone(geo.parse_mgrs("38SKC785335133"))  # odd number of digits
        self.assertIsNone(geo.parse_mgrs("hello"))

    def test_iraq_box_and_decimals(self):
        self.assertTrue(geo.in_iraq(33.3, 44.4))
        self.assertFalse(geo.in_iraq(44.4, 33.3))
        self.assertEqual(geo.parse_decimal("33,5"), 33.5)
        self.assertIsNone(geo.parse_decimal("abc"))


@tagged("post_install", "-at_install", "dma_fr")
class TestXlsx(BaseCase):
    def test_two_header_rows_and_moved_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "EORE"
        ws["A4"], ws["B4"], ws["C4"] = "Location *", None, "Date of report*"
        ws["A5"], ws["B5"] = "Province", "District"
        ws["A6"], ws["B6"], ws["C6"] = "Anbar", "Heet", datetime.datetime(2026, 9, 27)
        for r in range(7, 30):  # formatted but empty rows are ignored
            ws.cell(r, 1).value = None
        buf = io.BytesIO()
        wb.save(buf)
        sheet = xlsx.read_sheet(buf.getvalue(), "eore", 4, 2, 6)
        self.assertEqual([c.path for c in sheet.columns], ["Location * / Province", "Location * / District", "Date of report*"])
        self.assertEqual(sheet.find(["Location / District"])[0].letter, "B")
        self.assertEqual(sheet.find(["Date of report"])[0].letter, "C")
        self.assertEqual(sheet.rows, [(6, ["Anbar", "Heet", datetime.date(2026, 9, 27)])])


@tagged("post_install", "-at_install", "dma_fr")
class TestMaxml(BaseCase):
    def test_escaping_and_structure(self):
        rep = maxml.Report("form-guid", "DMA Form", "23", "FR-MRE-CI-1", "fr-guid", "Tester",
                           now=datetime.datetime(2026, 10, 1, 9, 0))
        rep.report_info(reporting_person="A & B <team>", reporting_date=datetime.date(2026, 9, 30))
        act = rep.activity("MREActivity", "MREIdentification", "obj-guid", "Education-CI-1")
        rep.cdf(act, "MRE.MULTI_SELECT.x", "g1", ["IED", "ERW"], "List")
        rep.cdf(act, "MRE.DATE.d", "g2", "2026-09-30", "Date")
        rep.point(act, "p", lat=33.5, lon=44.25)
        data = rep.tostring()
        self.assertTrue(data.startswith(b"\xff\xfe"))  # UTF-16 with BOM, as IMSMA got from the old tool
        root = etree.fromstring(data)
        self.assertEqual(root.findtext(".//ReportingPerson"), "A & B <team>")
        self.assertEqual([v.text for v in root.find(".//CDF").findall("Value")], ["IED", "ERW"])
        self.assertEqual(root.findtext(".//CDF[2]/Value"), "2026-09-30T15:00:00")
        self.assertEqual(root.findtext(".//XY/X"), "44.25000000")
        self.assertEqual(maxml.datatype_of("MRE.NUMBER.Bags"), "Numeric")


@tagged("post_install", "-at_install", "dma_fr")
class TestRefPackage(BaseCase):
    def test_package_round_trip_and_checksum(self):
        out = {name: [] for name in refdata.FILES}
        out["forms.csv"] = [{"guid": "g", "name": "F", "version": "1.0.01", "published_on": "", "active": "true"}]
        buf = io.BytesIO()
        refdata.write_package(out, {"source": "backup", "dump_date": "2025-03-02"}, buf)
        back, manifest = refdata.read_package(buf.getvalue())
        self.assertEqual(back["forms.csv"][0]["version"], "1.0.01")
        self.assertEqual(manifest["dump_date"], "2025-03-02")
        with self.assertRaises(ValueError):
            refdata.read_package(b"not a zip")
