"""Demo batches for the presentation captures (run in odoo-bin shell on dma_fr_livelike).

The operator sample's 4 rows, with people's names replaced by example names and the organisation shown as DMA.
Batch A stays with its problems (check / fix pages); batch B is fixed and exported (export page).
"""
import base64
import glob
import io
import os

from openpyxl import load_workbook

SAMPLE = r"C:/Users/Lenovo/Documents/imsma-importer/samples/eore_sep26"
NAMES = ["أحمد كريم (مثال)", "سارة علي (مثال)", "أحمد كريم (مثال)", "مصطفى حسن (مثال)"]

env = env(context=dict(env.context, lang="ar_001"))  # messages are written in the user's language
env.user.partner_id.name = "مسؤول البيانات (مثال)"
wb = load_workbook(glob.glob(SAMPLE + "/*.xlsm")[0], keep_vba=False)
ws = wb["EORE"]
for i, r in enumerate(range(6, 10)):
    ws.cell(r, 3).value = NAMES[i]          # C reported by
    ws.cell(r, 5).value = "DMA"             # E organisation
buf = io.BytesIO()
wb.save(buf)
data = buf.getvalue()
open(r"C:/Users/Lenovo/Documents/dma_field_reports/presentation/EORE_example.xlsx", "wb").write(data)


def make(label):
    atts = env["ir.attachment"].create([{"name": os.path.basename(p), "datas": base64.b64encode(open(p, "rb").read())}
                                        for p in sorted(glob.glob(SAMPLE + "/*.docx"))])
    b = env["dma.fr.batch"].create({
        "template_id": env.ref("dma_field_reports.template_eore").id,
        "office_id": env.ref("dma_field_reports.office_dma_hq").id,
        "source_file": base64.b64encode(data), "source_filename": "EORE_RMAC-N_2026-09.xlsx",
        "attachment_ids": [(6, 0, atts.ids)], "note": label})
    b.action_check()
    return b


# learned mappings from earlier test runs would hide the problems the pages show
env["dma.fr.value.map"].search([("origin", "=", "learned")]).unlink()
# locations "confirmed" by earlier test runs in this copy (not in IMSMA): back to unknown
env["dma.fr.ref.location"].with_context(active_test=False).search([("origin", "=", "confirmed")]).unlink()
env["dma.fr.batch"].search([]).unlink()  # this test copy only: start the screens from a clean list
a = make("presentation A")
b = make("presentation B")
W = env["dma.fr.fix.wizard"]
Opt = env["dma.fr.ref.option"]
FUND, MAC = "c0a8-cc01-160e6ac65d2-2fbf5fc1-1d-ab3", "c0a7-f58e-1773eacedb7-ef3a9ab4-10-dff"
haditha = env["dma.fr.ref.place"].search([("level", "=", "city"), ("name", "=", "Al Haditha")], limit=1)
for c in b.cell_ids.filtered(lambda c: c.state == "error"):
    if c.state != "error":
        continue
    if c.raw == "WRA":
        W.create({"cell_id": c.id, "option_id": Opt.search([("value", "=", "PMWRA"), ("cdf_ref_id.cdf_id", "=", FUND)]).id, "remember": False}).action_apply()
    elif c.raw == "NORAD":
        W.create({"cell_id": c.id, "option_id": Opt.search([("value", "=", "Norway"), ("cdf_ref_id.cdf_id", "=", FUND)]).id, "remember": False}).action_apply()
    elif c.field_id.cdf_guid == MAC:
        W.create({"cell_id": c.id, "option_id": Opt.search([("value", "=", "RMAC-N"), ("cdf_ref_id.cdf_id", "=", MAC)]).id, "remember": False}).action_apply()
for c in b.cell_ids.filtered(lambda c: c.code == "unknown_place"):
    if c.raw == "Alkasarah":
        loc = env["dma.fr.ref.location"].search([("localid", "=", "LOC-CI-40420")])
        W.create({"cell_id": c.id, "place_choice": "existing", "location_pick_id": loc.id, "apply_all": False}).action_apply()
    elif c.code == "unknown_place":
        W.create({"cell_id": c.id, "place_choice": "inside", "place_id": haditha.id, "apply_all": False}).action_apply()
b.invalidate_recordset()
print("B errors", b.error_count, "locations", b.locations_to_export)
b.action_export_locations()
b.action_confirm_locations()
b.action_export()
print("A", a.id, a.name, a.state, a.error_count, "| B", b.id, b.name, b.state)
env.cr.commit()
